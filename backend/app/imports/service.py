import csv
import hashlib
import io
import zipfile
from decimal import Decimal, InvalidOperation
from typing import Any, Literal
from uuid import UUID

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy.orm import Session

from app.audit import service as audit_service
from app.catalog import repository as catalog_repo
from app.idempotency import service as idempotency_service
from app.inventory import repository as inventory_repo
from app.shared.exceptions import bad_request
from app.shared.validation import INTEGER_MAX, reject_html

MAX_IMPORT_ROWS = 1000
MAX_IMPORT_BYTES = 2 * 1024 * 1024
MAX_XLSX_ARCHIVE_FILES = 128
MAX_XLSX_UNCOMPRESSED_BYTES = 16 * 1024 * 1024
MAX_XLSX_PART_BYTES = 8 * 1024 * 1024
MAX_XLSX_CELL_CHARACTERS = 2000
CatalogImportFormat = Literal["csv", "xlsx"]
ZIP_SIGNATURES = (b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")
TEMPLATE_COLUMNS = (
    "nombre",
    "sku",
    "precio",
    "costo",
    "categoria",
    "control_inventario",
    "stock_inicial",
    "umbral_stock",
    "codigo_barras",
)
# Excel recognizes the UTF-8 BOM and preserves Spanish accents when the owner
# opens the downloaded template. `_decode` already accepts the same BOM on
# upload through `utf-8-sig`.
TEMPLATE_CSV = "\ufeff" + ",".join(TEMPLATE_COLUMNS) + "\r\n"


def _decimal(value: str, *, label: str, required: bool) -> tuple[Decimal | None, str | None]:
    if not value:
        return (None, f"{label} es obligatorio") if required else (None, None)
    try:
        parsed = Decimal(value)
    except InvalidOperation:
        return None, f"{label} debe ser un número válido"
    if (
        not parsed.is_finite()
        or parsed < 0
        or parsed.as_tuple().exponent < -2
        or parsed >= Decimal("10000000000")
    ):
        return None, f"{label} debe ser positivo y tener máximo dos decimales"
    return parsed.quantize(Decimal("0.01")), None


def _integer(value: str, *, label: str) -> tuple[int | None, str | None]:
    if not value:
        return None, None
    try:
        parsed = int(value)
    except ValueError:
        return None, f"{label} debe ser un número entero"
    if parsed < 0:
        return None, f"{label} no puede ser negativo"
    if parsed > INTEGER_MAX:
        return None, f"{label} no puede superar {INTEGER_MAX}"
    return parsed, None


def _boolean(value: str) -> tuple[bool, str | None]:
    normalized = value.strip().lower()
    if normalized in {"", "no", "false", "0"}:
        return False, None
    if normalized in {"sí", "si", "true", "1"}:
        return True, None
    return False, "control_inventario debe ser sí o no"


def _decode(content: bytes) -> str:
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise bad_request("El CSV debe estar codificado en UTF-8") from exc


def _validate_headers(headers: list[str], *, source_label: str) -> None:
    if not headers:
        raise bad_request(f"El {source_label} no contiene encabezados")
    if len(headers) != len(set(headers)):
        raise bad_request(f"El {source_label} contiene encabezados repetidos")
    missing = [column for column in ("nombre", "precio") if column not in headers]
    unknown = [column for column in headers if column not in TEMPLATE_COLUMNS]
    if missing or unknown:
        details = []
        if missing:
            details.append(f"faltan columnas: {', '.join(missing)}")
        if unknown:
            details.append(f"columnas no reconocidas: {', '.join(unknown)}")
        raise bad_request("Encabezados inválidos; " + "; ".join(details))


def _parse_catalog_csv(content: bytes) -> list[tuple[int, dict[str, Any]]]:
    if not content.strip():
        raise bad_request("El archivo CSV está vacío")
    text = _decode(content)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    headers = [str(value or "").strip().lower() for value in (reader.fieldnames or [])]
    _validate_headers(headers, source_label="CSV")
    raw_rows: list[tuple[int, dict[str, Any]]] = []
    for row_number, raw in enumerate(reader, start=2):
        if not any(str(value or "").strip() for value in raw.values()):
            continue
        if len(raw_rows) >= MAX_IMPORT_ROWS:
            raise bad_request(f"El archivo admite máximo {MAX_IMPORT_ROWS} productos")
        raw_rows.append((row_number, dict(raw)))
    return raw_rows


def _validate_xlsx_package(content: bytes) -> None:
    if not content:
        raise bad_request("El archivo Excel está vacío")
    if len(content) > MAX_IMPORT_BYTES:
        raise bad_request("El archivo debe pesar 2 MB o menos")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_XLSX_ARCHIVE_FILES:
                raise bad_request("El archivo Excel contiene demasiados componentes")
            names = {entry.filename for entry in entries}
            if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                raise bad_request("El archivo no es un libro .xlsx válido")
            total_uncompressed = 0
            for entry in entries:
                path_parts = entry.filename.replace("\\", "/").split("/")
                if entry.filename.startswith(("/", "\\")) or ".." in path_parts:
                    raise bad_request("El archivo Excel contiene rutas no permitidas")
                if entry.flag_bits & 0x1:
                    raise bad_request("El archivo Excel está cifrado; quita la contraseña")
                if entry.file_size > MAX_XLSX_PART_BYTES:
                    raise bad_request("El archivo Excel excede los límites permitidos")
                total_uncompressed += entry.file_size
                if total_uncompressed > MAX_XLSX_UNCOMPRESSED_BYTES:
                    raise bad_request("El archivo Excel excede los límites permitidos")
            lowered_names = {name.casefold() for name in names}
            if any(name.endswith("vbaproject.bin") for name in lowered_names):
                raise bad_request("Los archivos con macros no están permitidos")
            content_types = archive.read("[Content_Types].xml").lower()
            if b"macroenabled" in content_types or b"vba" in content_types:
                raise bad_request("Los archivos con macros no están permitidos")
    except zipfile.BadZipFile as exc:
        raise bad_request(
            "No pudimos abrir el archivo Excel. Verifica que no esté cifrado ni dañado"
        ) from exc


def _parse_catalog_xlsx(content: bytes) -> list[tuple[int, dict[str, Any]]]:
    _validate_xlsx_package(content)
    try:
        workbook = load_workbook(
            io.BytesIO(content),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except (
        InvalidFileException,
        zipfile.BadZipFile,
        OSError,
        ValueError,
        KeyError,
        TypeError,
    ) as exc:
        raise bad_request(
            "No pudimos abrir el archivo Excel. Verifica que no esté cifrado ni dañado"
        ) from exc

    try:
        if len(workbook.worksheets) != 1:
            raise bad_request("El archivo Excel debe contener exactamente una hoja")
        worksheet = workbook.worksheets[0]
        if worksheet.max_row > MAX_IMPORT_ROWS + 1:
            raise bad_request(f"El archivo admite máximo {MAX_IMPORT_ROWS} productos")
        if worksheet.max_column > len(TEMPLATE_COLUMNS):
            raise bad_request(f"El archivo admite máximo {len(TEMPLATE_COLUMNS)} columnas")

        raw_rows: list[tuple[int, dict[str, Any]]] = []
        headers: list[str] | None = None
        for row_number, cells in enumerate(worksheet.iter_rows(), start=1):
            values: list[Any] = []
            for cell in cells:
                if cell.data_type == "f":
                    raise bad_request(
                        f"La celda {cell.coordinate} contiene una fórmula; reemplázala por su valor"
                    )
                value = cell.value
                if isinstance(value, str) and len(value) > MAX_XLSX_CELL_CHARACTERS:
                    raise bad_request(f"La celda {cell.coordinate} excede el límite de caracteres")
                values.append(value)
            if row_number == 1:
                headers = [str(value or "").strip().lower() for value in values]
                _validate_headers(headers, source_label="archivo Excel")
                continue
            if headers is None:
                raise bad_request("El archivo Excel no contiene encabezados")
            if not any(str(value or "").strip() for value in values):
                continue
            raw_rows.append((row_number, dict(zip(headers, values, strict=True))))
        return raw_rows
    finally:
        workbook.close()


def _validate_catalog_rows(
    db: Session,
    *,
    tenant_id: UUID,
    raw_rows: list[tuple[int, dict[str, Any]]],
) -> dict[str, Any]:

    existing_skus = {
        product.sku.upper()
        for product in catalog_repo.list_products(db, tenant_id=tenant_id, include_inactive=True)
        if product.sku
    }
    inactive_categories = {
        category.name.casefold()
        for category in catalog_repo.list_categories(db, tenant_id=tenant_id, include_inactive=True)
        if not category.is_active
    }
    existing_barcodes = {
        product.barcode
        for product in catalog_repo.list_products(db, tenant_id=tenant_id, include_inactive=True)
        if product.barcode
    }
    seen_barcodes: set[str] = set()
    seen_skus: set[str] = set()
    rows = []
    for row_number, raw in raw_rows:
        values = {
            str(key or "").strip().lower(): ("" if value is None else str(value).strip())
            for key, value in raw.items()
        }
        if not any(values.values()):
            continue
        errors: list[str] = []
        name = values.get("nombre", "").strip()
        if not name:
            errors.append("nombre es obligatorio")
        elif len(name) > 160:
            errors.append("nombre admite máximo 160 caracteres")
        else:
            try:
                reject_html(name)
            except ValueError as exc:
                errors.append(str(exc))
        sku = values.get("sku", "").upper() or None
        if sku and len(sku) > 100:
            errors.append("sku admite máximo 100 caracteres")
        if sku and sku in existing_skus:
            errors.append("sku ya existe en el catálogo")
        if sku and sku in seen_skus:
            errors.append("sku está repetido en el archivo")
        if sku:
            seen_skus.add(sku)
        price, price_error = _decimal(values.get("precio", ""), label="precio", required=True)
        cost, cost_error = _decimal(values.get("costo", ""), label="costo", required=False)
        if price_error:
            errors.append(price_error)
        if cost_error:
            errors.append(cost_error)
        category = values.get("categoria", "").strip() or None
        if category and len(category) > 120:
            errors.append("categoria admite máximo 120 caracteres")
        if category and category.casefold() in inactive_categories:
            errors.append("categoria existe pero está desactivada")
        track_inventory, inventory_error = _boolean(values.get("control_inventario", ""))
        if inventory_error:
            errors.append(inventory_error)
        initial_stock, stock_error = _integer(
            values.get("stock_inicial", ""), label="stock_inicial"
        )
        threshold, threshold_error = _integer(values.get("umbral_stock", ""), label="umbral_stock")
        if stock_error:
            errors.append(stock_error)
        if threshold_error:
            errors.append(threshold_error)
        if (initial_stock or 0) > 0 and not track_inventory:
            errors.append("activa control_inventario para cargar stock_inicial")
        if threshold is not None and not track_inventory:
            errors.append("activa control_inventario para definir umbral_stock")
        barcode = values.get("codigo_barras", "") or None
        if barcode and (len(barcode) > 100 or any(ord(c) < 33 or ord(c) > 126 for c in barcode)):
            errors.append("codigo_barras admite de 1 a 100 caracteres sin espacios")
        if barcode and (barcode in existing_barcodes or barcode in seen_barcodes):
            errors.append("codigo_barras ya existe o está repetido")
        if barcode:
            seen_barcodes.add(barcode)
        normalized = {
            "name": name,
            "sku": sku,
            "price_amount": str(price) if price is not None else None,
            "cost_price": str(cost) if cost is not None else None,
            "category_name": category,
            "track_inventory": track_inventory,
            "initial_stock": initial_stock or 0,
            "low_stock_threshold": threshold,
        }
        if barcode is not None:
            normalized["barcode"] = barcode
        rows.append(
            {
                "row_number": row_number,
                "status": "error" if errors else "valid",
                "normalized": normalized,
                "errors": errors,
            }
        )
    if not rows:
        raise bad_request("El archivo no contiene productos")
    error_rows = sum(1 for row in rows if row["errors"])
    return {
        "dry_run": True,
        "total_rows": len(rows),
        "valid_rows": len(rows) - error_rows,
        "error_rows": error_rows,
        "rows": rows,
        "created_products": 0,
        "created_categories": 0,
        "initial_stock_movements": 0,
    }


def validate_catalog_import(
    db: Session,
    *,
    tenant_id: UUID,
    content: bytes,
    file_format: CatalogImportFormat,
) -> dict[str, Any]:
    if file_format == "csv" and content.startswith(ZIP_SIGNATURES):
        raise bad_request("El contenido del archivo no coincide con el formato CSV seleccionado")
    raw_rows = _parse_catalog_csv(content) if file_format == "csv" else _parse_catalog_xlsx(content)
    return _validate_catalog_rows(db, tenant_id=tenant_id, raw_rows=raw_rows)


def validate_catalog_csv(db: Session, *, tenant_id: UUID, content: bytes) -> dict[str, Any]:
    """Backward-compatible CSV entry point used by existing callers/tests."""
    return validate_catalog_import(db, tenant_id=tenant_id, content=content, file_format="csv")


def _catalog_import_request_hash(
    content: bytes,
    *,
    file_format: CatalogImportFormat,
) -> str:
    # Preserve the historical CSV identity exactly. XLSX is namespaced so the
    # same bytes cannot replay an idempotent request created for another format.
    payload = content if file_format == "csv" else b"xlsx\0" + content
    return hashlib.sha256(payload).hexdigest()


def commit_catalog_import(
    db: Session,
    *,
    tenant_id: UUID,
    user_id: UUID,
    content: bytes,
    idempotency_key: str,
    file_format: CatalogImportFormat,
    commit: bool = True,
) -> tuple[int, dict[str, Any]]:
    request_hash = _catalog_import_request_hash(content, file_format=file_format)
    existing = idempotency_service.claim(
        db,
        tenant_id=tenant_id,
        key=idempotency_key,
        request_hash=request_hash,
    )
    if existing:
        return existing.response_status or 200, existing.response_body or {}
    from app.tenants.repository import lock_by_id

    lock_by_id(db, tenant_id)
    result = validate_catalog_import(
        db,
        tenant_id=tenant_id,
        content=content,
        file_format=file_format,
    )
    if result["error_rows"]:
        raise bad_request("Corrige todos los errores del archivo antes de confirmar")
    categories = {
        category.name.casefold(): category
        for category in catalog_repo.list_categories(db, tenant_id=tenant_id)
    }
    created_categories = 0
    initial_movements = 0
    for row in result["rows"]:
        values = row["normalized"]
        category_id = None
        category_name = values["category_name"]
        if category_name:
            category = categories.get(category_name.casefold())
            if not category:
                category = catalog_repo.create_category(
                    db,
                    tenant_id=tenant_id,
                    name=category_name,
                    description=None,
                    sort_order=0,
                )
                categories[category_name.casefold()] = category
                created_categories += 1
            category_id = category.id
        product = catalog_repo.create_product(
            db,
            tenant_id=tenant_id,
            category_id=category_id,
            name=values["name"],
            description=None,
            sku=values["sku"],
            barcode=values.get("barcode"),
            price_amount=Decimal(values["price_amount"]),
            cost_price=(Decimal(values["cost_price"]) if values["cost_price"] else None),
            track_inventory=values["track_inventory"],
            low_stock_threshold=values["low_stock_threshold"],
            image_position_x=50,
            image_position_y=50,
            image_zoom=1.0,
        )
        if values["initial_stock"] > 0:
            inventory_repo.create_movement(
                db,
                tenant_id=tenant_id,
                product_id=product.id,
                user_id=user_id,
                movement_type="adjustment",
                quantity_delta=values["initial_stock"],
                reason=f"Importación {file_format.upper()}: stock inicial",
            )
            initial_movements += 1
    response = {
        **result,
        "dry_run": False,
        "created_products": result["valid_rows"],
        "created_categories": created_categories,
        "initial_stock_movements": initial_movements,
    }
    audit_service.log(
        db,
        action="catalog.import",
        tenant_id=tenant_id,
        user_id=user_id,
        resource_type="catalog_import",
        resource_id=None,
        changes={
            key: response[key]
            for key in (
                "created_products",
                "created_categories",
                "initial_stock_movements",
            )
        }
        | {"file_format": file_format},
    )
    idempotency_service.store(
        db,
        tenant_id=tenant_id,
        key=idempotency_key,
        request_hash=request_hash,
        response_status=201,
        response_body=response,
    )
    if commit:
        db.commit()
    else:
        db.flush()
    return 201, response


def commit_catalog_csv(
    db: Session,
    *,
    tenant_id: UUID,
    user_id: UUID,
    content: bytes,
    idempotency_key: str,
) -> tuple[int, dict[str, Any]]:
    """Backward-compatible CSV entry point used by existing callers/tests."""
    return commit_catalog_import(
        db,
        tenant_id=tenant_id,
        user_id=user_id,
        content=content,
        idempotency_key=idempotency_key,
        file_format="csv",
    )
