"""Exercise the built native image's ABI, crypto, CA store and minimal surface.

No database or provider credentials are used. Database-backed request acceptance
remains a separate CI gate. The check runs as the image's configured uid 10001.
"""

import argparse
import subprocess

PROBE = r'''
import ctypes, hashlib, importlib, io, json, os, pathlib, platform, pwd
import importlib.util
import pyexpat, ssl, sys, uuid, zlib
from zoneinfo import ZoneInfo
assert os.getuid() == 10001 and os.getgid() == 10001
assert pwd.getpwuid(os.getuid()).pw_name == "kova"
assert sys.version_info[:2] == (3, 12)
assert sys.version_info >= (3, 12, 16)
assert pyexpat.version_info >= (2, 9, 0)
assert platform.machine() == "x86_64"
assert ssl.create_default_context().cert_store_stats()["x509_ca"] > 0
assert uuid.uuid4().version == 4 and uuid.uuid1().version == 1
assert ZoneInfo("America/Mexico_City").key == "America/Mexico_City"

stdlib = pathlib.Path(sys.base_prefix) / "lib/python3.12/lib-dynload"
extensions = []
for extension in sorted(stdlib.glob("*.so")):
    name = extension.name.split(".")[0]
    importlib.import_module(name)
    extensions.append(name)
for optional in ("readline", "_curses", "_curses_panel", "_uuid", "_tkinter"):
    assert importlib.util.find_spec(optional) is None, optional

import bcrypt, jwt, psycopg, psycopg.pq, sqlite3, uvloop
from cryptography.hazmat.primitives.ciphers.aead import AESGCM
from PIL import Image
assert psycopg.pq.__impl__ == "binary"
password = bcrypt.hashpw(b"runtime-probe", bcrypt.gensalt(rounds=4))
assert bcrypt.checkpw(b"runtime-probe", password)
key, nonce = AESGCM.generate_key(bit_length=256), os.urandom(12)
encrypted = AESGCM(key).encrypt(nonce, b"private", b"scope")
assert AESGCM(key).decrypt(nonce, encrypted, b"scope") == b"private"
jwt_key = "synthetic-runtime-test-key-32-bytes-minimum"
token = jwt.encode({"sub": "synthetic"}, jwt_key, algorithm="HS256")
assert jwt.decode(token, jwt_key, algorithms=["HS256"])["sub"] == "synthetic"
with sqlite3.connect(":memory:") as db:
    assert db.execute("select 1").fetchone() == (1,)
for format in ("PNG", "JPEG", "WEBP"):
    output = io.BytesIO()
    Image.new("RGB", (2, 2), "red").save(output, format=format)
    output.seek(0)
    Image.open(output).load()
loop = uvloop.new_event_loop()
loop.close()

# These business paths import their renderers lazily, after API startup.
from openpyxl import Workbook, load_workbook
from reportlab.pdfgen.canvas import Canvas
import qrcode
workbook, output = Workbook(), io.BytesIO()
workbook.active.append(["SKU", "Precio"])
workbook.save(output)
output.seek(0)
assert load_workbook(output, read_only=True).active.cell(1, 1).value == "SKU"
pdf = io.BytesIO()
canvas = Canvas(pdf)
canvas.drawString(20, 20, "Comprobante sintético")
canvas.save()
assert pdf.getvalue().startswith(b"%PDF-")
png = io.BytesIO()
qrcode.make("https://example.invalid/catalogo/prueba").save(png, format="PNG")
assert png.getvalue().startswith(b"\x89PNG")

# Resolve every shipped native extension and vendored shared library using the
# final runtime loader. A missing transitive ELF dependency fails this check.
shared_objects = []
for root in (pathlib.Path("/usr/local/lib"), pathlib.Path("/app/.venv/lib")):
    for library in sorted(root.rglob("*")):
        if library.is_file() and ".so" in library.name:
            ctypes.CDLL(str(library), mode=os.RTLD_LOCAL | os.RTLD_NOW)
            shared_objects.append(str(library))

for directory in ("/bin", "/sbin", "/usr/bin", "/usr/sbin", "/usr/local/bin", "/app/.venv/bin"):
    for tool in ("sh", "bash", "dash", "mount", "umount", "nsenter", "perl", "infocmp",
                 "systemctl", "apt", "dpkg", "uv", "pip", "pip3"):
        assert not pathlib.Path(directory, tool).exists(), (directory, tool)
assert not pathlib.Path("/app/app/tests").exists()
assert pathlib.Path("/var/lib/dpkg/status.d").is_dir()
provenance = json.loads(pathlib.Path("/usr/share/kova/native-runtime.json").read_text())
for package in provenance["debian_libraries"]:
    assert pathlib.Path("/", package["status"]).is_file(), package["name"]
    assert all(pathlib.Path(name).exists() for name in package["files"]), package["name"]
    for name, sha256 in package["sha256"].items():
        assert hashlib.sha256(pathlib.Path(name).read_bytes()).hexdigest() == sha256, name
assert provenance["embedded_expat"] == pyexpat.EXPAT_VERSION
assert provenance["zlib_runtime"] == zlib.ZLIB_RUNTIME_VERSION
from app.native import app, application
assert app is not application
schema = application.openapi()
assert "/health" in schema["paths"]
print(json.dumps({"uid": os.getuid(), "python": platform.python_version(),
                  "openssl": ssl.OPENSSL_VERSION,
                  "expat": pyexpat.EXPAT_VERSION, "zlib": zlib.ZLIB_RUNTIME_VERSION,
                  "ca_roots": ssl.create_default_context().cert_store_stats()["x509_ca"],
                  "stdlib_extensions": len(extensions), "shared_objects": len(shared_objects),
                  "psycopg": psycopg.pq.__impl__, "paths": len(schema["paths"])}))
'''


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--image", required=True)
    args = parser.parse_args()
    subprocess.run([
        "docker", "run", "--rm", "--platform", "linux/amd64", "--network", "none",
        "--read-only", "--cap-drop=ALL", "--security-opt=no-new-privileges",
        "--tmpfs", "/tmp:rw,noexec,nosuid,size=32m",
        "--env", "APP_ENV=ci",
        "--env", "SECRET_KEY=synthetic-runtime-test-key-32-bytes-minimum",
        "--env", "DATABASE_URL=postgresql+psycopg://pos@127.0.0.1:5432/synthetic",
        args.image, "/app/.venv/bin/python", "-c", PROBE,
    ], check=True)


if __name__ == "__main__":
    main()
